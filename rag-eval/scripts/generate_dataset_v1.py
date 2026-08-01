"""Generate self-authored, license-safe corpus assets for dataset-v1.0."""
from pathlib import Path
import hashlib, json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[1] / "dataset-v1.0"
RAW = ROOT / "corpus" / "raw"
CANONICAL = ROOT / "corpus" / "canonical-text"
MANIFEST = ROOT / "manifest" / "documents.jsonl"

DOMAINS = {
    "membership": ["租户生命周期规则", "成员邀请与加入规范", "角色权限矩阵", "成员状态校验", "管理员保留约束", "成员退出流程", "成员审计记录", "跨租户访问限制"],
    "file": ["文件上传准入规则", "原始与派生文件用途", "对象存储路径规范", "文件状态流转", "文件删除与保留", "租户级文件隔离", "文件类型白名单", "文件哈希与去重"],
    "document": ["文档创建规则", "文档版本快照", "文档状态管理", "异步解析流程", "解析失败处理", "派生文本产物", "当前版本切换", "文档访问边界"],
    "knowledge": ["章节优先分段策略", "固定窗口与重叠", "Embedding 模型调用", "Qdrant 索引结构", "Payload 租户过滤", "知识检索收敛", "Prompt 上下文约束", "问答引用规则"],
    "operations": ["JWT 认证边界", "RabbitMQ 消息约定", "Redis 连接治理", "OpenAI 密钥管理", "错误码响应规范", "关键日志规范", "超时与重试策略", "索引故障排查"],
}
FORMATS = ["md", "txt", "pdf", "docx", "md", "txt", "pdf", "docx"]

def text_for(domain, title, number):
    code = f"{domain.upper()}-{number:03d}"
    return f"""# {title}

## 文档目的

本文档定义 HexDoc 企业协作平台中“{title}”的统一规则。规则编号为 {code}，适用于生产、预发布和演示租户。任何自动化任务都必须记录规则编号、操作人和执行时间。

## 核心规则

1. 处理动作只有在所属租户状态为 ACTIVE 时才允许继续；租户被冻结时必须拒绝后续写入。
2. 规则 {code} 的默认执行窗口为每周三 02:17 至 02:47，最大重试次数为 {number % 4 + 2} 次；重试间隔不得少于 30 秒。
3. 每一次成功处理都必须写入 SUCCESS 状态，但 SUCCESS 只描述本规则的处理结果，不能推导成员、文件或文档的其他状态。

## 责任边界

调用方负责完成成员权限校验；底层资源服务只负责租户隔离、对象存储与元数据维护。系统生成的内容必须标记为 DERIVED，用户直接上传的内容标记为 ORIGINAL。

## 异常处理

当依赖服务不可用时，记录错误码 {code}-E1 并停止当前动作。不得用历史缓存结果替代实时权限判断。恢复后可按原规则编号重试，但必须保留第一次失败的审计记录。

## 示例

租户 hex-demo 在 02:20 触发 {title}。若租户为 ACTIVE 且操作人是有效成员，任务可继续；若成员状态不是 ACTIVE，应返回拒绝结果，而不是把任务标记为 SUCCESS。
"""

def write_docx(path, title, text):
    doc = Document(); sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
    normal = doc.styles['Normal']; normal.font.name = 'Hiragino Sans GB'; normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Hiragino Sans GB'); normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    for name, size, color in [('Heading 1', 16, '2E74B5'), ('Heading 2', 13, '2E74B5')]:
        style = doc.styles[name]; style.font.name = 'Hiragino Sans GB'; style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Hiragino Sans GB'); style.font.size = Pt(size); style.font.color.rgb = RGBColor.from_string(color)
    p = doc.add_paragraph(); r = p.add_run(title); r.font.name = 'Hiragino Sans GB'; r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Hiragino Sans GB'); r.bold = True; r.font.size = Pt(18); r.font.color.rgb = RGBColor(11, 37, 69)
    for line in text.splitlines()[1:]:
        if line.startswith('## '): doc.add_heading(line[3:], level=2)
        elif line and line[0].isdigit() and '. ' in line: doc.add_paragraph(line, style='List Number')
        elif line: doc.add_paragraph(line)
    doc.save(path)

def write_pdf(path, title, text):
    pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
    c = canvas.Canvas(str(path)); c.setFont('STSong-Light', 16); c.drawString(72, 760, title)
    c.setFont('STSong-Light', 10); y = 730
    for line in [x for x in text.splitlines() if x]:
        for start in range(0, len(line), 42):
            c.drawString(72, y, line[start:start+42]); y -= 16
            if y < 72: c.showPage(); c.setFont('STSong-Light', 10); y = 760
    c.save()

def main():
    RAW.mkdir(parents=True, exist_ok=True); CANONICAL.mkdir(parents=True, exist_ok=True)
    records=[]; number=1
    for domain, titles in DOMAINS.items():
        for idx, title in enumerate(titles):
            fmt=FORMATS[idx]; key=f"{domain}-{idx+1:02d}-v1"; text=text_for(domain, title, number); number += 1
            canonical = CANONICAL / f"{key}.txt"; canonical.write_text(text, encoding='utf-8')
            raw = RAW / f"{key}.{fmt}"
            if fmt in ('md','txt'): raw.write_text(text, encoding='utf-8')
            elif fmt == 'pdf': write_pdf(raw, title, text)
            else: write_docx(raw, title, text)
            records.append({"documentKey":key,"title":title,"format":fmt,"rawFile":f"corpus/raw/{raw.name}","canonicalTextFile":f"corpus/canonical-text/{canonical.name}","rawSha256":hashlib.sha256(raw.read_bytes()).hexdigest(),"domain":domain,"tags":[domain,"hexdoc","baseline-v1"],"language":"zh-CN"})
    MANIFEST.write_text('\n'.join(json.dumps(x,ensure_ascii=False) for x in records)+'\n',encoding='utf-8')

if __name__ == '__main__': main()
